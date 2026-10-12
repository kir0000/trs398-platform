"""Extract electron kQ tables (Rev.1 Tables 20 and 21) from the VBA module modRev1kQ of the electron worksheet.

The numbers live in VBA source (Select Case blocks), not on a sheet. Chamber wall data is read from the
hidden sheet "Tables". Everything is read-only: the workbook is never saved and no macro is run.

Usage: python importers/extract_electron_tables.py <electrons.xlsm> [out.json]
Requires oletools (pip install oletools) to read the VBA source.
"""
import json
import logging
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
logging.disable(logging.CRITICAL)


def vba_module(path, name):
    from oletools.olevba import VBA_Parser
    for _fn, _stream, vname, code in VBA_Parser(path).extract_macros():
        if vname.startswith(name):
            return code
    raise KeyError(f"VBA module {name} not found")


def parse_rows(code, func):
    body = code.split(f"Function {func}", 1)[1].split("End Function", 1)[0]
    rows = {}
    for m in re.finditer(r"Case (\d+): " + func + r" = Array\(([^)]*)\)", body):
        vals = [float(v.replace("#", "")) for v in m.group(2).split(",")]
        rows[int(m.group(1))] = [None if v < 0 else v for v in vals]   # -1 means no data
    return rows


def extract(path):
    code = vba_module(path, "modRev1kQ")
    grid = [float(v.replace("#", "")) for v in
            re.search(r"R50grid = Array\(([^)]*)\)", code).group(1).split(",")]
    names = dict((int(i), n) for i, n in re.findall(r'Case (\d+): kQ_chamberName = "([^"]*)"', code))
    t20, t21 = parse_rows(code, "T20row"), parse_rows(code, "T21row")
    ws = openpyxl.load_workbook(path, data_only=True)["Tables"]
    wall = {}
    for r in range(5, 21):
        wall[r - 4] = {"material": ws.cell(r, 20).value, "thickness": ws.cell(r, 21).value,
                       "plane_parallel": bool(ws.cell(r, 22).value)}
    chambers = [{"id": i, "name": names[i], "wall_material": wall[i]["material"],
                 "wall_thickness": wall[i]["thickness"], "plane_parallel": wall[i]["plane_parallel"],
                 "kq_table20": t20.get(i), "kq_table21": t21.get(i)} for i in sorted(names)]
    return {"tables_version": "TRS398-Rev1-2024", "r50_grid": grid, "chambers": chambers}


if __name__ == "__main__":
    data = extract(sys.argv[1])
    out = sys.argv[2] if len(sys.argv) > 2 else "core/data/electron_tables.json"
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"{len(data['chambers'])} chambers, {len(data['r50_grid'])} R50 nodes -> {out}")
