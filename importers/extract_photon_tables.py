"""Extract photon reference tables (Table 16, Table 10, chamber data) from the user's xlsm into JSON.

Read-only: the workbook is opened with openpyxl and never saved.
Usage: python importers/extract_photon_tables.py <path-to-photons.xlsm> [out.json]
"""
import json
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

SHEET = "Таблицы"
TPR_ROW = 4          # header row with TPR20,10 nodes in C..N
FIRST_CHAMBER_ROW = 5
LAST_CHAMBER_ROW = 30
TABLE10_FIRST, TABLE10_LAST = 36, 44


def extract(path):
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    tpr_nodes = [ws.cell(TPR_ROW, c).value for c in range(3, 15)]
    chambers = []
    for r in range(FIRST_CHAMBER_ROW, LAST_CHAMBER_ROW + 1):
        chambers.append({
            "id": ws.cell(r, 1).value,
            "name": ws.cell(r, 2).value,
            "kq": [ws.cell(r, c).value for c in range(3, 15)],
            "wall_material": ws.cell(r, 15).value,
            "wall_thickness": ws.cell(r, 16).value,
            "cavity_length_cm": ws.cell(r, 17).value,
            "source": ws.cell(r, 18).value,
        })
    table10 = []
    for r in range(TABLE10_FIRST, TABLE10_LAST + 1):
        table10.append({
            "v_ratio": ws.cell(r, 3).value,
            "pulsed": [ws.cell(r, c).value for c in (4, 5, 6)],
            "pulsed_scanned": [ws.cell(r, c).value for c in (7, 8, 9)],
        })
    return {"tables_version": "TRS398-Rev1-2024", "tpr_nodes": tpr_nodes,
            "chambers": chambers, "table10": table10}


if __name__ == "__main__":
    data = extract(sys.argv[1])
    out = sys.argv[2] if len(sys.argv) > 2 else "core/data/photon_tables.json"
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"{len(data['chambers'])} chambers, {len(data['tpr_nodes'])} TPR nodes, {len(data['table10'])} Table-10 rows -> {out}")
