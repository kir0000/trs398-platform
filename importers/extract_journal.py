"""Extract the journal (table "ЖурналТбл", sheet "Журнал" = "Journal" in the Russian worksheets) from a worksheet into JSON.

Read-only. Photon file: deviation from 1 cGy/MU; electron file: deviation from nominal. Both journals hold only
results (no raw readings), so they are imported into legacy_journal, not into measurement_sessions.

Usage: python importers/extract_journal.py <worksheet.xlsm> <out.json>
"""
import json
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")


def extract(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Журнал"]
    entries = []
    for r in range(7, ws.max_row + 1):
        date, dev, dw = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 8).value
        if date is None or dev in (None, "") or dw in (None, ""):
            continue
        entries.append({
            "date": date.date().isoformat(),
            "deviation_pct": float(dev),
            "note": (ws.cell(r, 5).value or "").strip() or None,
            "chamber": ws.cell(r, 6).value,
            "mq_nc_per_mu": ws.cell(r, 7).value,
            "dw_zmax": float(dw),
            "excel_signal": ws.cell(r, 4).value or None,   # cached SPC signal, used to validate core/spc.py
        })
    return entries


if __name__ == "__main__":
    data = extract(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"{len(data)} journal entries -> {sys.argv[2]}")
