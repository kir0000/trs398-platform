"""Export of a trend to CSV and Excel: the points with limits and signals, the summary, the baseline periods; both languages."""
import csv
import datetime as dt
import io

import openpyxl
import pytest

from app import xlsx
from tests import sqlcompat
from tests.sqlcompat import dict_row
from tests.test_app import client, post, q, world  # noqa: F401  (fixtures are discovered here)
from tests.test_compare_page import make_linac


@pytest.fixture(scope="module")
def beam(world, client):
    lid = make_linac(world, client, "Выгрузка", [("photon", 6, False), ("electron", 9, False)], shift_month=8)
    bid = q(client, "SELECT b.id FROM beams b WHERE b.linac_id=%s AND b.modality='photon'", lid)[0][0]
    return bid


def get(world_, who, beam_id, **params):
    return world_[who].get("/trends/export", params=dict(beam_id=beam_id, **params))


def test_excel_has_three_sheets_with_exact_values_and_limits(world, client, beam):
    r = get(world, "view", beam, fmt="xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Данные", "Сводка", "Базовые периоды"]
    ws = wb["Данные"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("Дата", "Отклонение, %", "MR, %") and len(rows) == 15                  # the header and 14 monthly points
    assert isinstance(rows[1][0], dt.datetime) and rows[1][0].date() == dt.date(2025, 7, 15)
    assert all(isinstance(r_[1], float) for r_ in rows[1:])
    head = list(rows[0])
    cl, lo, hi = head.index("CL, %"), head.index("−3σ, %"), head.index("+3σ, %")
    last = rows[-1]
    assert last[cl] is not None and last[lo] < last[cl] < last[hi]                                  # the limits of the series
    assert any(r_[head.index("Сигнал правил")] for r_ in rows[1:])                                  # the shift of a month is a signal
    assert ws.freeze_panes == "A2"
    summary = dict(wb["Сводка"].iter_rows(min_row=2, values_only=True))
    assert summary["Записей"] == 14 and "Выгрузка" in summary["Пучок"] and summary["Допуск (действие), %"] == 2.0
    assert summary["Пользователь"] and isinstance(summary["Выгружено"], dt.datetime)


def test_csv_russian_uses_semicolon_and_decimal_comma_english_the_other_way(world, client, beam):
    ru = get(world, "phys", beam, fmt="csv")
    assert ru.status_code == 200 and ru.headers["content-type"].startswith("text/csv") and ru.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(ru.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][0] == "Дата" and rows[1][0] == "15.07.2025" and len(rows) == 15
    assert "," in rows[1][1] and "." not in rows[1][1]                                               # 0,123456 (the Russian Excel reads it)
    world["phys"].get("/lang/en")
    try:
        en = get(world, "phys", beam, fmt="csv")
        rows = list(csv.reader(io.StringIO(en.content.decode("utf-8-sig")), delimiter=","))
        assert rows[0][:3] == ["Date", "Deviation, %", "MR, %"] and "." in rows[1][1] and "," not in rows[1][1]
        wb = openpyxl.load_workbook(io.BytesIO(get(world, "phys", beam, fmt="xlsx").content))
        assert wb.sheetnames == ["Data", "Summary", "Baseline periods"]
    finally:
        world["phys"].get("/lang/ru")


def test_baseline_periods_are_exported_and_points_before_them_have_no_limits(world, client, beam):
    uid = q(client, "SELECT id FROM users WHERE login='admin'")[0][0]
    with sqlcompat.connect(client.owner_url, row_factory=dict_row) as conn:
        conn.execute("INSERT INTO spc_baseline(beam_id, from_date, to_date, set_by) VALUES (%s, '2025-10-01', '2026-02-28', %s)", (beam, uid))
        conn.commit()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(world, "view", beam, fmt="xlsx").content))
        periods = list(wb["Базовые периоды"].iter_rows(values_only=True))
        assert periods[0][:3] == ("Период №", "С", "По") and len(periods) == 2 and periods[1][0] == 1
        data = list(wb["Данные"].iter_rows(values_only=True))
        head = list(data[0])
        first, last = data[1], data[-1]
        assert first[head.index("До первого базового периода")] == "да" and first[head.index("CL, %")] is None      # before the period: no limits
        assert last[head.index("Базовый период №")] == 1 and last[head.index("CL, %")] is not None
    finally:
        with sqlcompat.connect(client.owner_url, row_factory=dict_row) as conn:
            conn.execute("DELETE FROM spc_baseline WHERE beam_id = %s", (beam,))
            conn.commit()


def test_the_source_and_the_protocol_number_of_a_measured_point(world, client):
    from tests.test_app import create_measurement
    sid = int(create_measurement(world, "phys").headers["location"].rsplit("/", 1)[1])
    page = f"/measurements/{sid}"
    assert post(world["phys"], f"{page}/submit", page=page).status_code == 303
    assert post(world["rev"], f"{page}/approve", {"comment": "ok"}, page=page).status_code == 303
    no = q(client, "SELECT protocol_no FROM measurement_sessions WHERE id=%s", sid)[0][0]
    wb = openpyxl.load_workbook(io.BytesIO(get(world, "view", world["beam_id"], fmt="xlsx").content))
    data = list(wb["Данные"].iter_rows(values_only=True))
    head = list(data[0])
    assert data[1][head.index("Источник")] == "измерение" and data[1][head.index("Протокол №")] == no and data[1][head.index("Dw(zmax), Гр/MU")]


def test_errors_access_and_audit(world, client, beam):
    assert get(world, "view", beam, fmt="pdf").status_code == 422
    assert get(world, "view", 999999, fmt="xlsx").status_code == 404
    empty = q(client, "SELECT b.id FROM beams b WHERE b.modality='electron' AND b.id NOT IN (SELECT beam_id FROM legacy_journal)")
    assert not empty or get(world, "view", empty[0][0], fmt="csv").status_code == 404                  # a beam without data
    from fastapi.testclient import TestClient
    from app.main import app
    assert TestClient(app, follow_redirects=False).get("/trends/export", params=dict(beam_id=beam)).status_code == 303   # to the sign-in page
    assert q(client, "SELECT count(*) FROM audit_log WHERE action='trends_export' AND entity_id=%s", str(beam))[0][0] >= 2
    html = world["adm"].get(f"/trends?beam_id={beam}").text
    assert f"/trends/export?beam_id={beam}" in html and "Скачать Excel" in html and "Скачать CSV" in html


def test_the_xlsx_writer_escapes_and_names_the_sheets(tmp_path):
    data = xlsx.write([("A/B:C*?", [["x", "y"], ["<&>\x01", None], [float("nan"), True]]), ("A/B:C*?", [["z"]])])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["A B C", "A B C 2"]                                                      # forbidden characters, a repeated name
    assert [c.value for c in wb["A B C"][2]] == ["<&>", None]
