"""Export of the comparison of the energies of a linac to CSV and Excel: the same numbers as on the page, both languages."""
import csv
import datetime as dt
import io

import openpyxl
import pytest

from tests.test_app import client, q, world  # noqa: F401  (fixtures are discovered here)
from tests.test_compare_page import make_linac, plot_of

BEAMS = [("photon", 6, False), ("photon", 10, True), ("electron", 6, False), ("electron", 9, False), ("electron", 12, False)]


@pytest.fixture(scope="module")
def linac(world, client):
    return make_linac(world, client, "Выгрузка сравнения", BEAMS, shift_month=8)


def export(world_, who, linac_id, **params):
    return world_[who].get("/compare/export", params=dict(linac_id=linac_id, **params))


def test_excel_has_five_sheets_with_the_numbers_of_the_page(world, client, linac):
    r = export(world, "view", linac, fmt="xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Параметры", "Состояние", "По месяцам", "Корреляции", "Выводы"]
    params = dict(wb["Параметры"].iter_rows(min_row=2, values_only=True))
    assert params["Аппарат"] == "Выгрузка сравнения" and params["Тип пучков"] == "фотоны и электроны" and params["Допуск (действие), %"] == 2.0
    state = list(wb["Состояние"].iter_rows(values_only=True))
    assert state[0][:5] == ("Тип", "Пучок", "Записей", "Последняя", "Откл., %") and len(state) == 6          # five energies
    assert {(row[0], row[1]) for row in state[1:]} == {("γ", "6 MV"), ("γ", "10 MV FFF"), ("e⁻", "6 MeV"), ("e⁻", "9 MeV"), ("e⁻", "12 MeV")}
    assert all(row[2] == 14 and isinstance(row[3], dt.datetime) and isinstance(row[4], float) for row in state[1:])
    months = list(wb["По месяцам"].iter_rows(values_only=True))
    assert months[0][0] == "Месяц" and len(months) == 15 and months[1][0] == "2025-07"                           # 14 months
    head = list(months[0])
    shifted = next(r_ for r_ in months[1:] if r_[0] == "2026-03")
    assert shifted[head.index("Общий сдвиг")] == "вверх" and shifted[head.index("Вверх")] >= 4                  # the common shift of the page
    d = plot_of(world["view"].get(f"/compare?linac_id={linac}").text)                                          # the same matrix as the page plots
    first_beam = months[1][1]
    assert first_beam == pytest.approx(d["matrix"][0][0], abs=1e-3)
    corr = list(wb["Корреляции"].iter_rows(values_only=True))
    assert len(corr) == 6 and corr[1][1] == pytest.approx(1.0)
    assert len(list(wb["Выводы"].iter_rows(values_only=True))) >= 2


def test_filters_of_the_page_apply_to_the_export(world, client, linac):
    wb = openpyxl.load_workbook(io.BytesIO(export(world, "view", linac, fmt="xlsx", modality="electron").content))
    state = list(wb["Состояние"].iter_rows(values_only=True))
    assert len(state) == 4 and {row[0] for row in state[1:]} == {"e⁻"}
    params = dict(wb["Параметры"].iter_rows(min_row=2, values_only=True))
    assert params["Тип пучков"] == "только электроны"


def test_csv_in_both_languages(world, client, linac):
    ru = export(world, "phys", linac, fmt="csv")
    assert ru.status_code == 200 and ru.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(ru.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][:4] == ["Тип", "Пучок", "Записей", "Последняя"] and len(rows) == 6
    assert "," in rows[1][4] and "." not in rows[1][4] and rows[1][3].count(".") == 2                           # 0,123456 and 15.06.2026
    world["phys"].get("/lang/en")
    try:
        en = export(world, "phys", linac, fmt="csv")
        rows = list(csv.reader(io.StringIO(en.content.decode("utf-8-sig")), delimiter=","))
        assert rows[0][:4] == ["Type", "Beam", "Records", "Latest"] and "." in rows[1][4] and "," not in rows[1][4]
        wb = openpyxl.load_workbook(io.BytesIO(export(world, "phys", linac, fmt="xlsx").content))
        assert wb.sheetnames == ["Parameters", "State", "By month", "Correlations", "Conclusions"]
    finally:
        world["phys"].get("/lang/ru")


def test_errors_access_audit_and_the_buttons_of_the_page(world, client, linac):
    assert export(world, "view", linac, fmt="pdf").status_code == 422
    assert export(world, "view", linac, fmt="csv", modality="x").status_code == 422
    assert export(world, "view", 999999, fmt="csv").status_code == 404
    empty = q(client, "SELECT l.id FROM linacs l WHERE l.name NOT IN (SELECT DISTINCT linac FROM v_trend) LIMIT 1")
    assert not empty or export(world, "view", empty[0][0], fmt="csv").status_code == 404
    from fastapi.testclient import TestClient
    from app.main import app
    assert TestClient(app, follow_redirects=False).get("/compare/export", params=dict(linac_id=linac)).status_code == 303
    assert q(client, "SELECT count(*) FROM audit_log WHERE action='compare_export' AND entity_id=%s", str(linac))[0][0] >= 3
    html = world["adm"].get(f"/compare?linac_id={linac}").text
    assert f"/compare/export?linac_id={linac}" in html and "Скачать Excel" in html and "Скачать CSV" in html
